from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db.models import Q
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .models import Product, Unit

SAMPLE_HEADERS = ["Name", "Category", "Unit", "Stock Qty", "Cost Price", "Selling Price"]
SAMPLE_ROWS = [
    ["Khaman", "Sweets", "kg", 10, 180, 240],
    ["Samosa", "Farsan", "pc", 50, 4, 8],
]


def _units_for_business(business):
    return Unit.objects.filter(Q(business=business) | Q(business__isnull=True)).order_by("abbreviation")


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        ws.column_dimensions[col_cells[0].column_letter].width = max(10, length + 2)


def build_sample_workbook(business) -> bytes:
    """An .xlsx an Admin can download, fill in, and re-upload — includes a
    second sheet listing every unit abbreviation this business can use, so
    imports don't fail on a typo'd unit."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(SAMPLE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in SAMPLE_ROWS:
        ws.append(row)

    units_ws = wb.create_sheet("Valid Units")
    units_ws.append(["Abbreviation", "Name"])
    for cell in units_ws[1]:
        cell.font = Font(bold=True)
    for unit in _units_for_business(business):
        units_ws.append([unit.abbreviation, unit.name])

    _autosize(ws)
    _autosize(units_ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportResult:
    def __init__(self):
        self.created = 0
        self.updated = 0
        self.errors = []  # list of (row_number, message)

    @property
    def has_errors(self):
        return bool(self.errors)

    @property
    def total_ok(self):
        return self.created + self.updated


def import_products_from_workbook(business, file_obj) -> ImportResult:
    result = ImportResult()

    try:
        wb = load_workbook(file_obj, data_only=True)
    except Exception:
        result.errors.append((0, "Couldn't read that file — make sure it's a valid .xlsx Excel file."))
        return result

    ws = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]
    valid_units = {unit.abbreviation.lower(): unit for unit in _units_for_business(business)}

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(cell in (None, "") for cell in row):
            continue  # blank row, skip silently

        name, category, unit_abbr, stock_qty, cost_price, selling_price = (list(row) + [None] * 6)[:6]

        name = str(name).strip() if name is not None else ""
        if not name:
            result.errors.append((row_num, "Missing product name."))
            continue

        unit_abbr = str(unit_abbr).strip() if unit_abbr is not None else ""
        unit = valid_units.get(unit_abbr.lower())
        if unit is None:
            result.errors.append(
                (row_num, f"Unknown unit '{unit_abbr}' — check the 'Valid Units' sheet for options.")
            )
            continue

        try:
            stock_qty = Decimal(str(stock_qty)) if stock_qty not in (None, "") else Decimal("0")
            cost_price = Decimal(str(cost_price))
            selling_price = Decimal(str(selling_price))
        except InvalidOperation:
            result.errors.append((row_num, "Stock qty, cost price and selling price must all be numbers."))
            continue

        category = str(category).strip() if category is not None else ""

        existing = Product.objects.for_business(business).filter(name__iexact=name).first()
        if existing:
            existing.category = category
            existing.unit = unit
            existing.stock_qty = stock_qty
            existing.cost_price = cost_price
            existing.selling_price = selling_price
            existing.is_active = True
            existing.save()
            result.updated += 1
        else:
            Product.objects.create(
                business=business,
                name=name,
                category=category,
                unit=unit,
                stock_qty=stock_qty,
                cost_price=cost_price,
                selling_price=selling_price,
            )
            result.created += 1

    return result
